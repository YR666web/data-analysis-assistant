"""离线构建内部知识库向量索引（TDD §9.2）。

加载 knowledge_base/ 语料 → 分块 → BGE 向量化 → 持久化到 vector_store/。
"""
from __future__ import annotations

import os
import shutil
import sys
from pathlib import Path

# 允许以 `python scripts/build_knowledge_base.py` 直接运行时导入项目内模块
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from dotenv import load_dotenv
from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

from utils.logger_utils import get_logger

# 加载项目根目录的 .env
_PROJECT_ROOT = Path(__file__).resolve().parent.parent
load_dotenv(_PROJECT_ROOT / ".env")

logger = get_logger("data_analysis_assistant.build")

# 与 tools/knowledge_base_tool.py 保持一致
_COLLECTION_NAME = "data_analysis_kb"
_CHUNK_SIZE = 500
_CHUNK_OVERLAP = 50


def _resolve_path(p: str) -> Path:
    """把 .env 里的（可能是相对的）路径解析为绝对路径。"""
    path = Path(p)
    if not path.is_absolute():
        path = _PROJECT_ROOT / path
    return path.resolve()


def _get_embedding() -> HuggingFaceEmbeddings:
    """按 TDD §12 读取 Embedding 模型：优先本地路径，否则用模型名。"""
    model_path = os.getenv("EMBEDDING_MODEL_PATH", "").strip()
    if model_path:
        model_name = str(_resolve_path(model_path))
    else:
        model_name = os.getenv("EMBEDDING_MODEL", "BAAI/bge-base-zh-v1.5")
    return HuggingFaceEmbeddings(model_name=model_name)


# ---- 多格式文档解析（TDD §2：python-docx / pypdf / openpyxl / pandas）----

def _load_text(path: Path) -> str:
    """读取文本文件，自动兼容 UTF-8 与 GB18030（覆盖 GBK/GB2312）。"""
    raw = path.read_bytes()
    for enc in ("utf-8", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _load_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs)


def _load_xlsx(path: Path) -> str:
    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None)
    parts = [f"[Sheet: {name}]\n{df.to_csv(index=False)}" for name, df in sheets.items()]
    return "\n\n".join(parts)


_LOADERS = {
    ".txt": _load_text,
    ".md": _load_text,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".xlsx": _load_xlsx,
    ".xls": _load_xlsx,
}


def load_documents(kb_dir: Path) -> list[Document]:
    """递归加载 kb_dir 下的所有支持文档，来源元数据用文件名。"""
    docs: list[Document] = []
    for path in sorted(kb_dir.rglob("*")):
        if not path.is_file():
            continue
        loader = _LOADERS.get(path.suffix.lower())
        if loader is None:
            logger.warning("跳过不支持的类型: %s", path.name)
            continue
        try:
            text = loader(path)
        except Exception as e:  # noqa: BLE001
            logger.error("解析失败 %s: %s", path.name, e)
            continue
        if not text.strip():
            logger.warning("跳过空内容: %s", path.name)
            continue
        docs.append(Document(page_content=text, metadata={"source": path.name}))
        logger.info("加载 %s（%d 字符）", path.name, len(text))
    return docs


def build_index() -> int:
    """构建向量索引，返回分块总数。"""
    kb_dir = _resolve_path(os.getenv("KNOWLEDGE_BASE_DIR", "./knowledge_base"))
    vs_dir = _resolve_path(os.getenv("VECTOR_STORE_DIR", "./vector_store"))

    if not kb_dir.exists():
        logger.error("语料目录不存在: %s", kb_dir)
        return 0

    logger.info("语料目录: %s", kb_dir)
    documents = load_documents(kb_dir)
    if not documents:
        logger.error("未加载到任何文档，构建终止")
        return 0

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=_CHUNK_SIZE,
        chunk_overlap=_CHUNK_OVERLAP,
    )
    chunks = splitter.split_documents(documents)
    logger.info("分块完成：%d 个文档 → %d 个分块", len(documents), len(chunks))

    embedding = _get_embedding()

    # 清空旧索引，实现「可重建」
    if vs_dir.exists():
        shutil.rmtree(vs_dir)
    vs_dir.mkdir(parents=True, exist_ok=True)

    vector_store = Chroma.from_documents(
        chunks,
        embedding,
        persist_directory=str(vs_dir),
        collection_name=_COLLECTION_NAME,
    )
    # 显式关闭 chromadb 系统，强制把 HNSW 索引 .bin 文件刷到磁盘。
    # chromadb 的 Rust 后端若在进程退出前不 stop，.bin 文件可能未落盘，
    # 导致后续检索时报 "Error loading hnsw index"。
    vector_store._client._system.stop()
    logger.info("向量索引已持久化到: %s", vs_dir)
    return len(chunks)


if __name__ == "__main__":
    count = build_index()
    logger.info("构建完成，共 %d 个分块", count)
