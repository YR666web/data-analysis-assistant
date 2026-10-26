import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pandas as pd
from langchain_core.messages import AIMessage, HumanMessage, ToolMessage

from api import context
from api.monitor import MonitorCallbackHandler, tool_outcome
from tools.table_tools import analyze_table, inspect_table
from tools.upload_file_read_tool import list_session_files, read_file_content


class TableAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        context.set_session("table_test", str(self.root))
        self.frame = pd.DataFrame({"日期": ["2026-10-19"] * 7 + ["2026-10-26"],
                                   "分类": ["抗生素", "非抗生素"] * 4,
                                   "销售额": ["0.1", "0.2", "0.3", "0.4", "0.5", "0.6", "100", "999"]})
        self.frame.to_excel(self.root / "sales.xlsx", index=False)

    def tearDown(self):
        context.reset_session()
        self.temp.cleanup()

    def analyze(self, **kwargs):
        return analyze_table.invoke({"filename": "sales.xlsx", **kwargs})

    def test_weekly_total_includes_rows_beyond_preview_excludes_next_week(self):
        result = self.analyze(date_column="日期", start_date="2026-10-19", end_date="2026-10-25",
                              metrics=[{"column": "销售额", "op": "sum", "alias": "销售总额"}])
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual(result["matched_rows"], 7)
        self.assertEqual(result["results"], [{"销售总额": "102.1"}])
        self.assertEqual(result["date_coverage"]["max"], "2026-10-26")

    def test_non_antibiotic_exact_filter(self):
        result = self.analyze(filters=[{"column": "分类", "op": "eq", "value": "非抗生素"}])
        self.assertEqual(result["matched_rows"], 4)
        self.assertTrue(all(row["分类"] == "非抗生素" for row in result["results"]))

    def test_grouping_and_pagination_do_not_change_aggregate(self):
        result = self.analyze(group_by=["分类"], metrics=[{"column": "销售额", "op": "sum", "alias": "总额"}], limit=1)
        self.assertEqual(result["total_result_rows"], 2)
        self.assertEqual(result["next_offset"], 1)
        self.assertEqual(result["results"][0]["总额"], "100.9")

    def test_no_data_reports_coverage_not_zero_sales(self):
        result = self.analyze(date_column="日期", start_date="2027-01-01", end_date="2027-01-07",
                              metrics=[{"column": "销售额", "op": "sum", "alias": "总额"}])
        self.assertEqual(result["status"], "no_data")
        self.assertEqual(result["results"], [])
        self.assertEqual(result["date_coverage"]["min"], "2026-10-19")

    def test_multiple_sheets_need_explicit_choice(self):
        with pd.ExcelWriter(self.root / "multi.xlsx") as writer:
            self.frame.to_excel(writer, sheet_name="销售", index=False)
            self.frame.head(1).to_excel(writer, sheet_name="说明", index=False)
        inspected = inspect_table.invoke({"filename": "multi.xlsx"})
        self.assertEqual([s["rows"] for s in inspected["sheets"]], [8, 1])
        self.assertEqual(analyze_table.invoke({"filename": "multi.xlsx"})["status"], "error")
        self.assertEqual(analyze_table.invoke({"filename": "multi.xlsx", "sheet": "销售"})["matched_rows"], 8)

    def test_bad_numeric_value_fails_instead_of_silently_dropping(self):
        self.frame.loc[7, "销售额"] = "未知"
        self.frame.to_excel(self.root / "sales.xlsx", index=False)
        result = self.analyze(metrics=[{"column": "销售额", "op": "sum", "alias": "总额"}])
        self.assertEqual(result["status"], "error")

    def test_missing_values_reported(self):
        self.frame.loc[7, "销售额"] = ""
        self.frame.to_excel(self.root / "sales.xlsx", index=False)
        result = self.analyze(metrics=[{"column": "销售额", "op": "sum", "alias": "总额"}])
        self.assertEqual(result["results"][0]["总额"], "102.1")
        self.assertTrue(result["warnings"])

    def test_invalid_column_operator_and_date_are_errors(self):
        for kwargs in [{"group_by": ["不存在"]}, {"filters": [{"column": "分类", "op": "eval"}]},
                       {"start_date": "2026-10-19"}, {"date_column": "销售额"}]:
            with self.subTest(kwargs=kwargs):
                self.assertEqual(self.analyze(**kwargs)["status"], "error")

    def test_csv_duplicate_header_and_full_read(self):
        (self.root / "dup.csv").write_text("a,a\n1,2\n", encoding="utf-8")
        self.assertEqual(inspect_table.invoke({"filename": "dup.csv"})["status"], "error")
        self.frame.to_csv(self.root / "sales.csv", index=False)
        result = analyze_table.invoke({"filename": "sales.csv", "offset": 5})
        self.assertEqual(len(result["results"]), 3)

    def test_session_manifest_and_traversal_protection(self):
        self.assertEqual(list_session_files.invoke({})["files"][0]["filename"], "sales.xlsx")
        for filename in ["../sales.xlsx", str(self.root / "sales.xlsx")]:
            self.assertEqual(read_file_content.invoke({"filename": filename})["status"], "error")

    def test_text_paging_and_word_tables(self):
        (self.root / "long.txt").write_text("正文" * 12000, encoding="utf-8")
        first = read_file_content.invoke({"filename": "long.txt"})
        second = read_file_content.invoke({"filename": "long.txt", "offset": first["next_offset"]})
        self.assertEqual(first["content"] + second["content"], "正文" * 12000)
        self.assertIsNone(second["next_offset"])
        from docx import Document
        doc = Document()
        doc.add_paragraph("销售说明")
        doc.add_table(rows=1, cols=1).cell(0, 0).text = "表格金额 123"
        doc.save(self.root / "sales.docx")
        self.assertIn("表格金额 123", read_file_content.invoke({"filename": "sales.docx"})["content"])

    def test_failed_tool_arguments_are_not_a_report(self):
        from agent.main_agent import _extract_report
        messages = [HumanMessage(content="生成报告"), AIMessage(content="", tool_calls=[
            {"id": "x", "name": "generate_markdown", "args": {"content": "未保存的假报告"}}]),
            ToolMessage(content="错误：磁盘不可写", tool_call_id="x")]
        self.assertEqual(_extract_report(self.root, messages, 0), "")


class ToolOutcomeTests(unittest.IsolatedAsyncioTestCase):
    async def test_real_agent_graph_can_call_session_tools(self):
        from agent import main_agent
        from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
        from api.monitor import manager
        class LocalModel(FakeMessagesListChatModel):
            def bind_tools(self, tools, **kwargs):
                return self
        model = LocalModel(responses=[
            AIMessage(content="", tool_calls=[{"id": "inspect", "name": "inspect_table", "args": {"filename": "sales.csv"}}]),
            AIMessage(content="", tool_calls=[{"id": "sum", "name": "analyze_table", "args": {
                "filename": "sales.csv", "metrics": [{"column": "amount", "op": "sum", "alias": "total"}]}}]),
            AIMessage(content="", tool_calls=[{"id": "write", "name": "generate_markdown", "args": {
                "filename": "report.md", "content": "# Sales\nTotal: 0.3\nSource: sales.csv"}}]),
            AIMessage(content="Done")])
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "sales.csv").write_text("amount\n0.1\n0.2\n", encoding="utf-8")
            context.set_session("graph_test", directory)
            manager.clear_streams("graph_test")
            try:
                with patch.object(main_agent, "get_llm", return_value=model):
                    result = await main_agent.main_agent_run("Read sales.csv and generate a sales report", 1)
                self.assertIn("Total: 0.3", result["report"])
                events = [event for event in manager.event_snapshot("graph_test") if event["event"] == "tool_result"]
                self.assertEqual(len(events), 3)
                self.assertTrue(all(event["data"]["success"] for event in events))
            finally:
                context.reset_session()
                manager.clear_streams("graph_test")

    async def test_attachment_to_weekly_report_pipeline_without_external_model(self):
        from agent import main_agent
        from tools.markdown_tools import generate_markdown
        from utils import path_utils
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            upload, output = root / "upload", root / "output"
            upload.mkdir()
            pd.DataFrame({"日期": ["2026-10-19"] * 6 + ["2026-10-26"],
                          "销售额": [10, 20, 30, 40, 50, 60, 999]}).to_excel(upload / "sales.xlsx", index=False)
            captured = {}
            class LocalAgent:
                async def ainvoke(self, inputs, config):
                    captured.update(inputs)
                    result = analyze_table.invoke({"filename": "sales.xlsx", "date_column": "日期",
                        "start_date": "2026-10-19", "end_date": "2026-10-25",
                        "metrics": [{"column": "销售额", "op": "sum", "alias": "总额"}]})
                    if result["status"] != "ok":
                        raise AssertionError(result)
                    body = "# 销售周报\n日期：2026-10-19 至 2026-10-25\n来源：sales.xlsx\n销售额：" + result["results"][0]["总额"]
                    generate_markdown.invoke({"filename": "周报.md", "content": body})
                    return {"messages": [AIMessage(content="文件已生成")]}
            with patch.object(path_utils, "get_session_dir", return_value=output), \
                 patch.object(path_utils, "get_upload_dir", return_value=upload), \
                 patch.object(main_agent, "create_deep_agent", return_value=LocalAgent()), \
                 patch.object(main_agent, "soft_review", new_callable=AsyncMock, return_value=None):
                try:
                    result = await main_agent.run_deep_agent("根据上传表格生成销售周报", "weekly_test")
                    self.assertIn("销售额：210", result)
                    self.assertNotIn("999", result)
                    self.assertEqual((output / "周报.md").read_text(encoding="utf-8"), result)
                    self.assertIn("sales.xlsx", captured["messages"][0].content)
                finally:
                    context.reset_session()

    async def test_failure_and_no_data_events(self):
        callback = MonitorCallbackHandler("test")
        for output, status in [("PDF 转换失败：错误", "error"),
                               ({"status": "error", "error": "bad"}, "error"),
                               (ToolMessage(content='{"status":"no_data"}', tool_call_id="a"), "no_data")]:
            self.assertEqual(tool_outcome(output), status)
            callback._tool_names["a"] = "test_tool"
            with patch("api.monitor.monitor.report", new_callable=AsyncMock) as report:
                await callback.on_tool_end(output, run_id="a")
                self.assertEqual(report.call_args.args[3]["success"], status != "error")


if __name__ == "__main__":
    unittest.main()
